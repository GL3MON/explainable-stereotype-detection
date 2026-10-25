"""Build an editable review deck using the supplied BTP template.

Run collect_deck_data.py first. Requires python-pptx; outputs remain local.
"""
from copy import deepcopy
import json
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FOLDER = HERE.parent
DATA = json.loads((HERE / 'deck_data.json').read_text())
PRS = Presentation(FOLDER / 'BTP Presentation Template-3.pptx')
PROTOTYPES = list(PRS.slides)
R = DATA['results']
NAVY = '25354A'
BLUE = '4472C4'
GREEN = '4F8567'
TEXT = '263343'
GREY = '637082'
LINE = 'DCE4EC'
PALE_BLUE = 'EEF3FB'
PALE_GREEN = 'EFF6F0'
PALE_AMBER = 'FBF6E9'
PALE_PURPLE = 'F3EFF8'
WHITE = 'FFFFFF'
FONT = 'Calibri'
WIDTH = PRS.slide_width / 914400
HEIGHT = PRS.slide_height / 914400
SLIDE_LOG = []

REFS = [
    ('K. S. Shejole and P. Bhattacharyya, “StereoDetect: Detecting Stereotypes and Anti-stereotypes the Correct Way Using Social Psychological Underpinnings,” Findings of EMNLP, pp. 4051–4082, 2025.', 'https://aclanthology.org/2025.findings-emnlp.216/'),
    ('T. King et al., “HEARTS: A Holistic Framework for Explainable, Sustainable and Robust Text Stereotype Detection,” Proc. IJCNLP–AACL, pp. 1–18, 2025.', 'https://aclanthology.org/2025.ijcnlp-long.1/'),
    ('Z. Lan et al., “ALBERT: A Lite BERT for Self-supervised Learning of Language Representations,” ICLR, 2020.', 'https://arxiv.org/abs/1909.11942'),
    ('B. Warner et al., “Smarter, Better, Faster, Longer: A Modern Bidirectional Encoder for Fast, Memory Efficient, and Long Context Finetuning and Inference,” arXiv:2412.13663, 2024.', 'https://arxiv.org/abs/2412.13663'),
    ('W. Wang et al., “MiniLM: Deep Self-Attention Distillation for Task-Agnostic Compression of Pre-Trained Transformers,” NeurIPS, 2020.', 'https://arxiv.org/abs/2002.10957'),
    ('Sentence Transformers, “all-MiniLM-L6-v2,” model card. Accessed Oct. 22, 2026.', 'https://huggingface.co/sentence-transformers/all-MiniLM-L6-v2'),
    ('Y. Zhang et al., “Qwen3 Embedding: Advancing Text Embedding and Reranking Through Foundation Models,” arXiv:2506.05176, 2025.', 'https://arxiv.org/abs/2506.05176'),
    ('Qwen Team, “Qwen3.5-2B,” model card and model configuration. Accessed Oct. 22, 2026.', 'https://huggingface.co/Qwen/Qwen3.5-2B'),
    ('E. J. Hu et al., “LoRA: Low-Rank Adaptation of Large Language Models,” ICLR, 2022.', 'https://arxiv.org/abs/2106.09685'),
    ('M. T. Ribeiro, S. Singh, and C. Guestrin, “Why Should I Trust You?: Explaining the Predictions of Any Classifier,” KDD, pp. 1135–1144, 2016.', 'https://arxiv.org/abs/1602.04938'),
    ('S. M. Lundberg and S.-I. Lee, “A Unified Approach to Interpreting Model Predictions,” NeurIPS, 2017.', 'https://arxiv.org/abs/1705.07874'),
]


def rgb(value):
    return RGBColor.from_string(value)


def text(slide, x, y, w, h, content, size=22, color=TEXT, bold=False,
         align=PP_ALIGN.LEFT, font=FONT, margin=0, valign=MSO_ANCHOR.TOP):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(content.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(5)
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
    return shape


def box(slide, x, y, w, h, fill=PALE_BLUE, border=LINE):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.07
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(border)
    shape.line.width = Pt(.8)
    return shape


def card(slide, x, y, w, h, title, body, fill=PALE_BLUE, size=21):
    box(slide, x, y, w, h, fill)
    text(slide, x+.2, y+.16, w-.4, .45, title, size=22, bold=True, color=NAVY)
    text(slide, x+.2, y+.79, w-.4, h-.92, body, size=size)


def note(slide, content, sources=()):
    slide.notes_slide.notes_text_frame.text = content + (
        '\n\nSources:\n' + '\n'.join(str(source) for source in sources) if sources else '')


def source_line(slide, value):
    text(slide, 1.02, 6.58, 11.15, .2, value, size=10, color=GREY)


def rewrite(shape, value, size=None, color=None):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = value
    if size is not None:
        p.font.size = Pt(size)
        p.font.name = FONT
    if color is not None:
        p.font.color.rgb = rgb(color)


def slide(section, subtitle, prototype=1, sources=(), speaker=''):
    original = PROTOTYPES[prototype]
    new = PRS.slides.add_slide(original.slide_layout)
    for shape in list(new.shapes):
        new.shapes._spTree.remove(shape.element)
    for shape in original.shapes:
        if shape.has_text_frame and shape.name.startswith('Content Placeholder'):
            continue
        element = deepcopy(shape.element)
        for node in element.iter():
            for attr, value in list(node.attrib.items()):
                if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                    rel = original.part.rels[value]
                    node.set(attr, new.part.relate_to(rel.target_part if not rel.is_external else rel.target_ref,
                                                     rel.reltype, is_external=rel.is_external))
        new.shapes._spTree.insert_element_before(element, 'p:extLst')
    number = len(SLIDE_LOG)+1
    for shape in new.shapes:
        if shape.name.startswith('Title') and shape.has_text_frame:
            rewrite(shape, section)
        if shape.name.startswith('Slide Number'):
            rewrite(shape, str(number), size=16)
    if subtitle:
        text(new, 1.02, 1.64, 11.1, .3, subtitle, size=18, color=BLUE)
    SLIDE_LOG.append({'number': number, 'heading': section, 'subtitle': subtitle})
    note(new, speaker, sources)
    return new


def arrow(slide, x1, y1, x2, y2, color=GREY):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(1.5)
    end = OxmlElement('a:tailEnd')
    end.set('type', 'triangle')
    shape.line._get_or_add_ln().append(end)
    return shape


def flow(slide, labels, y=2.55, colors=None, x=1.03, width=11.26, h=1.16, size=19, connect=True):
    gap = .32
    w = (width-gap*(len(labels)-1))/len(labels)
    positions = []
    for i, label in enumerate(labels):
        left = x+i*(w+gap)
        box(slide, left, y, w, h, (colors or [PALE_BLUE]*len(labels))[i])
        text(slide, left+.08, y+.12, w-.16, h-.2, label, size=size,
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i and connect:
            arrow(slide, left-gap+.025, y+h/2, left-.025, y+h/2)
        positions.append((left, y, w, h))
    return positions


def table(slide, headers, values, widths, y=2.12, row_height=.52, size=18, highlight=()):
    graphic = slide.shapes.add_table(len(values)+1, len(headers), Inches(1.02), Inches(y),
                                     Inches(sum(widths)), Inches(row_height*(len(values)+1)))
    result = graphic.table
    for i, width in enumerate(widths):
        result.columns[i].width = Inches(width)
    for r, row in enumerate([headers, *values]):
        for c, value in enumerate(row):
            cell = result.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(.12)
            cell.margin_top = cell.margin_bottom = Inches(.07)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(PALE_BLUE if r == 0 else PALE_GREEN if r in highlight else WHITE if r%2 == 0 else 'F8FAFC')
            cell.text = str(value)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(size)
                p.font.bold = r == 0
                p.font.color.rgb = rgb(NAVY if r == 0 else TEXT)
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
    return result


def compare_chart(slide, categories, base, tuned, x=1.04, y=2.2, w=7.3, h=3.65):
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series('Base', [v*100 for v in base])
    chart_data.add_series('Fine-tuned', [v*100 for v in tuned])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                  Inches(w), Inches(h), chart_data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.legend.font.name = FONT
    chart.legend.font.size = Pt(16)
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 100
    chart.value_axis.major_unit = 25
    chart.value_axis.tick_labels.font.size = Pt(12)
    chart.value_axis.tick_labels.number_format = '0"%"'
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = rgb('E7ECF1')
    chart.category_axis.tick_labels.font.name = FONT
    chart.category_axis.tick_labels.font.size = Pt(17)
    chart.category_axis.has_major_gridlines = False
    chart.plots[0].gap_width = 100
    for series, color in zip(chart.series, ('BCCCE8', '76A98B')):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(color)
        series.format.line.fill.background()
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.number_format = '0.00"%"'
    plot.data_labels.font.name = FONT
    plot.data_labels.font.size = Pt(15)
    plot.data_labels.font.color.rgb = rgb(TEXT)
    return chart


def metric(slide, x, y, w, value, caption, fill=PALE_GREEN):
    box(slide, x, y, w, 1.35, fill)
    text(slide, x+.15, y+.17, w-.3, .6, value, size=34, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    text(slide, x+.12, y+.87, w-.24, .3, caption, size=16, align=PP_ALIGN.CENTER)


def percent(value):
    return f'{value*100:.2f}%'


# 1. Title. Keep the institution, logo, geometry, and editable identity fields.
s = slide('', '', prototype=0)
for shape in s.shapes:
    if not shape.has_text_frame:
        continue
    if shape.name.startswith('Title'):
        rewrite(shape, 'StereoDetect\nExplainable Stereotype Detection', size=36)
    elif shape.name.startswith('Subtitle'):
        if 'Guided' in shape.text:
            rewrite(shape, 'Guided By,\n[Guide name]', size=22)
        else:
            rewrite(shape, 'BY\n[Student names and roll numbers]', size=24)
SLIDE_LOG[-1] = {'number': 1, 'heading': 'BTP Title', 'subtitle': 'Editable title and identity details'}
note(s, 'The project title is a working title. The presenter will add the approved title, student names, roll numbers, and guide name. The supplied template logo and institution footer are retained.')

# 2. Introduction.
s = slide('Introduction', 'Five labels capture different forms of group-related language', speaker='Introduce StereoDetect as a benchmark-based project. The dataset and definitions are prior work; this project compares model families, task adaptation, structured output, and diagnostic explanations.', sources=[REFS[0][1], 'data/README.md'])
card(s, 1.02, 2.08, 5.46, 1.86, 'Why this task matters', 'Group references and generalizations require context-sensitive classification.', size=21)
card(s, 6.72, 2.08, 5.57, 1.86, 'Project focus', 'Compare model predictions and inspect the words influencing their decisions.', PALE_GREEN, 21)
names = ['Anti-stereotype', 'Stereotype', 'Neutral without\ntarget term', 'Neutral with\ntarget term', 'Bias']
flow(s, [f'{i}\n{name}' for i, name in enumerate(names)], y=4.32, h=1.43, size=18,
     colors=[PALE_BLUE, PALE_AMBER, PALE_GREEN, PALE_PURPLE, 'F8EFEF'], connect=False)
source_line(s, 'Dataset definitions: StereoDetect [1]. Labels match the supplied CSV annotations.')

# 3–4. Literature review.
s = slide('Literature Review', 'Dataset foundations and explainable stereotype detection', prototype=2, sources=[REFS[0][1], REFS[1][1]], speaker='StereoDetect supplies the five-class task and benchmark definitions. HEARTS supplies a related ALBERT-based binary stereotype detection and explainability approach. Its task and dataset differ from this project; its published scores are not mixed into the local comparison.')
card(s, 1.02, 2.12, 5.46, 3.44, 'StereoDetect · Shejole & Bhattacharyya', 'Definition-aligned benchmark for stereotypes, anti-stereotypes, neutrality, and bias.\n\nProvides the five-class labels used in this project. [1]', size=22)
card(s, 6.72, 2.12, 5.57, 3.44, 'HEARTS · King et al.', 'Combines ALBERT-based stereotype detection with LIME/SHAP explanations.\n\nMotivates diagnostic word attribution for our five-class setting. [2]', PALE_GREEN, 22)
source_line(s, 'The HEARTS binary task and StereoDetect five-class task use different evaluation settings.')

s = slide('Literature Review', 'Model families used in the comparison', prototype=2,
          sources=[REFS[i][1] for i in (2, 3, 4, 5, 6, 7, 8)],
          speaker='This slide describes public model backbones. The exact classifier fitting details for the MiniLM and Qwen3-Embedding baseline runs are not present in the supplied code; their scores are reported from DL-Review-2.docx. Later diagrams mark those task heads as schematic.')
table(s, ['Model family', 'Core idea', 'Role in this project'], [
    ['MiniLM [5], [6]', 'Distilled transformer embeddings', 'Compact baseline'],
    ['Qwen3-Embedding [7]', 'Decoder-derived text embeddings', 'Embedding baseline'],
    ['ALBERT [3]', 'Factorized embeddings; shared layers', 'Sentence classifier'],
    ['ModernBERT-large [4]', 'Local/global encoder attention', 'Sentence classifier'],
    ['Qwen3.5-2B + LoRA [8], [9]', 'Hybrid decoder; task adaptation', 'Label or JSON output'],
], [3.15, 4.25, 3.85], row_height=.62, size=18)
source_line(s, 'Architectures: original papers and official model cards [3]–[9].')

# 5. Motivation.
s = slide('Motivation', 'Test both classification quality and the usefulness of the output', prototype=3,
          sources=[REFS[0][1], *DATA['sources'].values()], speaker='Frame these as project questions, not as claims that no earlier work exists. Dataset-specific label meanings make task adaptation worth measuring. The richer-input label task and sentence-only JSON task must be interpreted separately.')
card(s, 1.02, 2.16, 3.6, 3.65, '1 · Adaptation', 'How much does task-specific fine-tuning improve the existing Qwen checkpoint?', size=23)
card(s, 4.85, 2.16, 3.6, 3.65, '2 · Output scope', 'Can one sentence produce the category, target group, and label together?', PALE_GREEN, 23)
card(s, 8.68, 2.16, 3.6, 3.65, '3 · Explanation', 'Which words influence an encoder decision, including a wrong prediction?', PALE_PURPLE, 23)
source_line(s, 'Project questions evaluated on the supplied StereoDetect splits.')

# 6. Problem statement.
s = slide('Problem Statement', 'Two Qwen task contracts, evaluated on the same held-out test split', prototype=4,
          sources=[DATA['sources']['base_labels'], DATA['sources']['ft_labels'], DATA['sources']['base_json'], DATA['sources']['ft_json']],
          speaker='The label task is conditioned on Category, Target, and Sentence. The structured task only sees Sentence and predicts all three output fields. Category and target are never fed as gold inputs in JSON generation. All model evaluation uses the existing 1,738-example test split. Base denotes the named Qwen checkpoint without a task-specific LoRA adapter.')
card(s, 1.02, 2.12, 5.46, 2.95, 'Task A · Label prediction', 'Input: Category + Target + Sentence\n\nOutput: one integer label, 0–4\n\nDecision: highest-scoring label token', size=22)
card(s, 6.72, 2.12, 5.57, 2.95, 'Task B · All-field prediction', 'Input: Sentence only\n\nOutput: category + target + label JSON\n\nDecision: greedy JSON generation', PALE_GREEN, 22)
box(s, 1.02, 5.34, 11.27, .76, PALE_AMBER)
text(s, 1.22, 5.5, 10.87, .4, 'Compare base and fine-tuned runs within each task; their input information differs.', size=21, bold=True)
source_line(s, 'Primary metrics: label accuracy and macro F1. Category/target accuracy is assessed separately.')

# 7. Overall system.
s = slide('Architecture', 'Overall experimental pipeline', prototype=5, sources=['src/stereodetect/', 'data/README.md'], speaker='The dataset is split before model training. Validation selects checkpoints; held-out test evaluation reports task-specific scores. Explanations are produced for saved ALBERT and ModernBERT-large classifiers. Public-backbone baseline scores are retained from the review document.')
flow(s, ['StereoDetect\nCSV splits', 'Task-specific\ninputs', 'Model training\n/ adaptation', 'Held-out\npredictions', 'Metrics +\nexplanations'], y=2.37, h=1.2, size=19, colors=[PALE_BLUE, PALE_GREEN, PALE_BLUE, PALE_GREEN, PALE_PURPLE])
card(s, 1.02, 4.04, 3.6, 1.81, 'Data', 'Train 6,080 · validation 868\nTest 1,738', size=20)
card(s, 4.85, 4.04, 3.6, 1.81, 'Outputs', 'Five labels, or structured\ncategory/target/label JSON', PALE_GREEN, 20)
card(s, 8.68, 4.04, 3.6, 1.81, 'Inspection', 'Class-wise scores and\nLIME/SHAP diagnostics', PALE_PURPLE, 20)
source_line(s, 'Implementation: stereodetect package. Dataset split sizes are counted from local CSV files.')

# 8. MiniLM.
s = slide('Architecture', 'all-MiniLM-L6-v2 · Compact sentence-embedding backbone', prototype=5,
          sources=[REFS[4][1], REFS[5][1], DATA['review_source']], speaker='The published all-MiniLM-L6-v2 backbone uses six transformer layers, 384-dimensional hidden states, mean pooling, and normalized sentence embeddings. The dashed task stage is schematic: the supplied review records its five-class score, but does not include the downstream classifier implementation. Do not imply that a particular classifier or training recipe was verified.')
flow(s, ['Sentence', 'Tokenizer', 'MiniLM encoder\n6 layers', 'Mean pooling\n+ normalize', '384-D\nembedding'], y=2.45, h=1.22, size=19)
arrow(s, 11.29, 3.72, 11.29, 4.58)
arrow(s, 11.29, 4.58, 8.96, 4.58)
box(s, 4.42, 4.16, 4.5, .85, PALE_GREEN)
text(s, 4.57, 4.33, 4.2, .45, 'Task classifier → five labels', size=22, bold=True, align=PP_ALIGN.CENTER)
text(s, 1.15, 5.36, 10.98, .7, 'A compact sentence representation for the review baseline.\nThe embedding backbone is shown above; the classification stage is schematic.', size=21, align=PP_ALIGN.CENTER)
source_line(s, 'Backbone: MiniLM and official all-MiniLM-L6-v2 model card [5], [6]. Baseline result: review document.')

# 9. Qwen embedding.
s = slide('Architecture', 'Qwen3-Embedding-0.6B · Decoder-derived sentence representation', prototype=5,
          sources=[REFS[6][1], 'https://huggingface.co/Qwen/Qwen3-Embedding-0.6B', DATA['review_source']], speaker='Qwen3-Embedding-0.6B uses a dense Qwen3 decoder backbone with 28 layers. Its official model extracts a final-token representation and supports an embedding dimension up to 1,024. The supplied materials report a five-class baseline score but omit the downstream classifier details, so that stage is explicitly schematic.')
flow(s, ['Sentence', 'Tokenizer', 'Qwen3 decoder\n28 layers', 'Final-token\npooling', 'Up to 1024-D\nembedding'], y=2.45, h=1.22, size=19)
arrow(s, 11.29, 3.72, 11.29, 4.58)
arrow(s, 11.29, 4.58, 8.96, 4.58)
box(s, 4.42, 4.16, 4.5, .85, PALE_GREEN)
text(s, 4.57, 4.33, 4.2, .45, 'Task classifier → five labels', size=22, bold=True, align=PP_ALIGN.CENTER)
text(s, 1.15, 5.36, 10.98, .7, 'A larger sentence representation for the review baseline.\nThe embedding backbone is shown above; the classification stage is schematic.', size=21, align=PP_ALIGN.CENTER)
source_line(s, 'Backbone: Qwen3 Embedding paper and official 0.6B model card [7]. Baseline result: review document.')

# 10. ALBERT.
s = slide('Architecture', 'ALBERT-base-v2 · Sentence-only five-class classifier', prototype=5,
          sources=[REFS[2][1], 'src/stereodetect/training/albert.py', 'outputs/albert_stereodetect/model/config.json'], speaker='ALBERT factorizes token embeddings and shares parameters across repeated transformer layers. The local classifier consumes Sentence only, applies the ALBERT pooler to the classification token, and produces five label logits. Category and Target are metadata, not model inputs. This differs from the wording in one row of the supplied review document; the diagram follows the actual code and run configuration.')
flow(s, ['Sentence\n+ special tokens', 'Factorized\nembeddings', 'Shared ALBERT\nencoder ×12', '[CLS] pooler\n+ classifier', 'Five label\nprobabilities'], y=2.52, h=1.3, size=19)
card(s, 1.02, 4.25, 5.46, 1.72, 'Parameter efficiency', 'Embedding projection and cross-layer sharing reduce learned parameters.', size=21)
card(s, 6.72, 4.25, 5.57, 1.72, 'Local experiment', 'Six epochs · effective batch 64\nLearning rate 2 × 10⁻⁵', PALE_GREEN, 21)
source_line(s, 'ALBERT [3]. Local task and head: saved configuration and training implementation.')

# 11. ModernBERT-large, with the actual saved mean-pooling head.
s = slide('Architecture', 'ModernBERT-large · Sentence-only five-class classifier', prototype=5,
          sources=[REFS[3][1], 'outputs/modernbert_large_lr3e5_warmup30_wd001_ep3_seed3407/model/config.json', 'src/stereodetect/training/modernbert.py'], speaker='Only the large variant is used throughout this deck. The saved configuration has 28 encoder layers, hidden width 1,024, local attention width 128, and global attention every third layer. The classification pooler is mean pooling in the saved configuration; it is not a CLS-only head. RoPE and GeGLU are public backbone design features. The training settings shown are for the additional saved large-model run, rather than an inferred recipe for the earlier review score.')
flow(s, ['Sentence\n+ tokenizer', 'Token\nembeddings', 'Local/global\nencoder ×28', 'Mean pooling\n+ classifier', 'Five label\nprobabilities'], y=2.45, h=1.22, size=19)
box(s, 2.75, 4.0, 7.85, .8, PALE_PURPLE)
text(s, 2.95, 4.18, 7.45, .4, 'RoPE · GeGLU · periodic global attention', size=23, bold=True, align=PP_ALIGN.CENTER)
text(s, 1.1, 5.19, 11.1, .85, 'Saved large-model run: 3 epochs · effective batch 32 · LR 3 × 10⁻⁵\nCheckpoint selected by validation macro F1.', size=22, align=PP_ALIGN.CENTER)
source_line(s, 'ModernBERT [4]. Large variant: 28 layers, hidden width 1024; details verified in the saved model config.')

# 12. Qwen3.5 hybrid language backbone and task heads.
s = slide('Architecture', 'Qwen3.5-2B · Shared language backbone, two task adapters', prototype=5,
          sources=[REFS[7][1], 'https://huggingface.co/Qwen/Qwen3.5-2B/raw/main/config.json', REFS[8][1], 'src/stereodetect/models/qwen/'], speaker='This project uses only the language path of the conditional-generation model. The official 2B text configuration contains 24 layers arranged in six repetitions of three linear-attention layers followed by one full-attention layer. Gated DeltaNet is used for the linear-attention path. LoRA is trained separately for the label and JSON tasks; the adapters are not interchangeable. Rank 16, alpha 32, dropout 0.05 are taken from the implementation. No vision input is used.')
flow(s, ['Task prompt\n+ tokenizer', 'Hybrid language\ndecoder ×24', 'Language-model\nhead'], y=2.22, h=1.06, width=8.3, size=20)
box(s, 9.62, 2.22, 2.65, 1.06, PALE_PURPLE)
text(s, 9.76, 2.4, 2.37, .68, 'Task LoRA\nr=16 · α=32', size=20, bold=True, align=PP_ALIGN.CENTER)
arrow(s, 10.94, 3.3, 10.94, 3.46, BLUE)
arrow(s, 10.94, 3.46, 5.18, 3.46, BLUE)
arrow(s, 5.18, 3.46, 5.18, 3.3, BLUE)
box(s, 1.02, 3.63, 11.25, .7, PALE_BLUE)
text(s, 1.2, 3.79, 10.89, .35, 'Repeated block: 3 × gated linear attention + 1 × full attention', size=22, bold=True, align=PP_ALIGN.CENTER)
card(s, 1.02, 4.64, 5.46, 1.48, 'Label task', 'Category + Target + Sentence → label', PALE_GREEN, 20)
card(s, 6.72, 4.64, 5.57, 1.48, 'JSON task', 'Sentence → category + target + label', PALE_GREEN, 20)
source_line(s, 'Qwen3.5-2B model configuration [8]; LoRA [9]. Text pathway only; separate task adapters.')

# 13. Explainability pipeline.
s = slide('Architecture', 'Aligned word-level explanations for the encoder classifiers', prototype=5,
          sources=[REFS[1][1], REFS[9][1], REFS[10][1], 'src/stereodetect/explanations/encoder.py'], speaker='Both explainers treat each word occurrence as a feature and replace absent words with the tokenizer mask token. LIME uses 1,000 perturbations. Kernel SHAP uses an all-masked baseline and up to 512 coalitions with no L1 feature selection. Predictions are made on the same perturbed strings by the saved encoder classifier. Explanations concern class probabilities, not moral judgments about words or calibrated prediction confidence. The 100-example sample is deliberately diagnostic and not representative test statistics.')
flow(s, ['Sentence\nword positions', 'Masked-word\nvariants', 'Saved encoder\nprobabilities', 'LIME /\nKernel SHAP', 'Word scores\n+ masking checks'], y=2.44, h=1.22, size=19, colors=[PALE_BLUE, PALE_GREEN, PALE_BLUE, PALE_PURPLE, PALE_GREEN])
card(s, 1.02, 4.17, 5.46, 1.8, 'Matched features', 'Same word positions, masking token, and explained class for both methods.', size=21)
card(s, 6.72, 4.17, 5.57, 1.8, 'Diagnostic interpretation', 'Agreement describes model sensitivity; it does not establish that a prediction is correct.', PALE_AMBER, 21)
source_line(s, 'Method sources: HEARTS [2], LIME [10], SHAP [11]. Local explanation settings are recorded per run.')

# 14. Experimental setup.
s = slide('Experimental Results', 'Dataset, training setup, and evaluation protocol', prototype=6,
          sources=['data/train.csv', 'data/val.csv', 'data/test.csv', 'outputs/albert_stereodetect/run_config.json', 'outputs/modernbert_large_lr3e5_warmup30_wd001_ep3_seed3407/run_config.json', 'src/stereodetect/training/qwen_labels.py'], speaker='All local comparisons use the same stored train/validation/test splits. The test file hash matches the new base benchmark. Training configurations below are for runs that have saved artifacts, including the additional ModernBERT-large run. Earlier review-only embedding baselines have no supplied fitting recipe. Fine-tuned checkpoints are selected using validation data. Base runs use vLLM; fine-tuned standalone evaluations use Transformers/PEFT. Runtime and throughput are therefore not compared as controlled measurements.')
for x, label, value, fill in [(1.02, 'Training examples', '6,080', PALE_BLUE), (4.85, 'Validation examples', '868', PALE_GREEN), (8.68, 'Test examples', '1,738', PALE_PURPLE)]:
    metric(s, x, 2.1, 3.6, value, label, fill)
table(s, ['Saved training run', 'Epochs', 'Effective batch', 'Learning rate'], [
    ['ALBERT-base-v2', '6', '64', '2 × 10⁻⁵'],
    ['ModernBERT-large', '3', '32', '3 × 10⁻⁵'],
    ['Qwen3.5-2B LoRA (each task)', '3', '16', '2 × 10⁻⁴'],
], [4.1, 1.65, 2.7, 2.8], y=3.87, row_height=.5, size=18)
source_line(s, 'Checkpoint selection: validation macro F1. Test metrics: accuracy, macro F1, per-class F1, and JSON field accuracy.')

# 15. Cross-model results, preserving the review run and saved large run separately.
s = slide('Experimental Results', 'Model comparison on the 1,738-example test split', prototype=6,
          sources=[DATA['review_source'], *DATA['sources'].values()], speaker='MiniLM, Qwen3-Embedding and the first ModernBERT-large row are the rounded figures in DL-Review-2.docx. The user identifies that review model as ModernBERT-large. The second ModernBERT-large row is a different, independently saved run, not a replacement for the review score. ALBERT and Qwen values are from standalone saved evaluations. The highest numerical score uses Category and Target as inputs, so this is a descriptive table and not an input-matched model ranking. All percentages are rounded to two decimals.')
values = [
    ['all-MiniLM-L6-v2', 'Review baseline', '77.33', '74.29'],
    ['Qwen3-Embedding-0.6B', 'Review baseline', '91.08', '91.11'],
    ['ALBERT-base-v2', 'Sentence only', percent(R['albert']['accuracy'])[:-1], percent(R['albert']['macro_f1'])[:-1]],
    ['ModernBERT-large · review', 'Review run', '92.98', '93.05'],
    ['ModernBERT-large · saved', 'Sentence only', percent(R['modernbert']['accuracy'])[:-1], percent(R['modernbert']['macro_f1'])[:-1]],
    ['Qwen3.5-2B · label LoRA', 'Category + target + sentence', percent(R['ft_labels']['accuracy'])[:-1], percent(R['ft_labels']['macro_f1'])[:-1]],
    ['Qwen3.5-2B · JSON LoRA', 'Sentence only', percent(R['ft_json']['accuracy'])[:-1], percent(R['ft_json']['macro_f1'])[:-1]],
]
table(s, ['Model / run', 'Input / source', 'Accuracy (%)', 'Macro F1 (%)'], values,
      [4.1, 3.6, 1.75, 1.8], y=2.05, row_height=.47, size=17, highlight=(6, 7))
source_line(s, 'Review baselines and saved evaluations are distinguished. Qwen label and JSON tasks use different input information.')

# 16. Base vs task adaptation, labels.
base, tuned = R['base_labels'], R['ft_labels']
s = slide('Experimental Results', 'Qwen3.5-2B · Base versus fine-tuned label prediction', prototype=6,
          sources=[DATA['sources']['base_labels'], DATA['sources']['ft_labels']], speaker='Both runs use Category, Target and Sentence, and choose among the five numeric label-token logits. The base benchmark evaluates the Qwen/Qwen3.5-2B checkpoint without a task adapter; the fine-tuned run loads outputs/qwen35_2b_stereodetect/adapter. Accuracy gain is computed in percentage points, not relative percent. Engines differ (vLLM versus Transformers/PEFT), so this is an observed task-adaptation comparison rather than an isolated engine-controlled ablation.')
compare_chart(s, ['Accuracy', 'Macro F1'], [base['accuracy'], base['macro_f1']], [tuned['accuracy'], tuned['macro_f1']])
metric(s, 8.73, 2.32, 3.53, f'+{(tuned["accuracy"]-base["accuracy"])*100:.2f} pp', 'Observed accuracy gain')
card(s, 8.73, 3.97, 3.53, 2.12, 'Same task inputs', 'Category + Target\n+ Sentence\nAll 1,738 test examples', PALE_BLUE, 20)
source_line(s, 'Base: 583 correct. Fine-tuned: 1,660 correct. Base = named Qwen checkpoint without the task adapter.')

# 17. Base vs task adaptation, JSON label scoring.
base, tuned = R['base_json'], R['ft_json']
s = slide('Experimental Results', 'Qwen3.5-2B · Base versus fine-tuned all-field generation', prototype=6,
          sources=[DATA['sources']['base_json'], DATA['sources']['ft_json']], speaker='Both runs receive only Sentence and generate category, target and label JSON. This chart scores the label field only; it does not mean all fields are correct. Greedy generation and a 96-token cap are used in the saved task configuration. The base run has 599 correct labels, the adapter run 1,627. Both saved evaluations record zero invalid labels. The distinct complete-JSON result appears on the next slide.')
compare_chart(s, ['Label accuracy', 'Label macro F1'], [base['accuracy'], base['macro_f1']], [tuned['accuracy'], tuned['macro_f1']])
metric(s, 8.73, 2.32, 3.53, f'+{(tuned["accuracy"]-base["accuracy"])*100:.2f} pp', 'Observed label accuracy gain')
card(s, 8.73, 3.97, 3.53, 1.98, 'Sentence-only input', 'Generate all three fields.\n\nThis chart scores the label.', PALE_BLUE, 20)
source_line(s, 'Base: 599 correct labels. Fine-tuned: 1,627 correct labels. Both use the complete 1,738-example test split.')

# 18. Complete structured-output scores, independently recomputed for the base run.
base_fields, tuned_fields = DATA['fields']['base_json'], DATA['fields']['ft_json']
s = slide('Experimental Results', 'Structured output · Category, target, label, and complete JSON', prototype=6,
          sources=[base_fields['source'], tuned_fields['source'], 'outputs/eval/json_test/field_validation.json'], speaker='Every saved completion is aligned with its source test row. Category is compared after trimming, case folding and collapsing whitespace. Target uses exact string comparison, including capitalization and underscores. Label must be an integer matching the gold label. Complete JSON accuracy requires all three fields to be correct for the same example. Both sets contain valid JSON objects for all 1,738 rows, so low base accuracy is semantic rather than just JSON formatting. The fine-tuned all-field count agrees with the existing field_validation.json. Base field metrics were computed from saved completions while preparing this deck; no model was rerun.')
table(s, ['Metric', 'Base (%)', 'Fine-tuned (%)', 'Fine-tuned correct'], [
    ['Category accuracy', f'{base_fields["category"]*100:.2f}', f'{tuned_fields["category"]*100:.2f}', '1,727 / 1,738'],
    ['Target exact match', f'{base_fields["target"]*100:.2f}', f'{tuned_fields["target"]*100:.2f}', '1,706 / 1,738'],
    ['Label accuracy', f'{base_fields["label"]*100:.2f}', f'{tuned_fields["label"]*100:.2f}', '1,627 / 1,738'],
    ['All three fields correct', f'{base_fields["all_fields"]*100:.2f}', f'{tuned_fields["all_fields"]*100:.2f}', '1,614 / 1,738'],
], [3.6, 2.0, 2.55, 3.1], y=2.23, row_height=.64, size=20, highlight=(4,))
text(s, 1.2, 5.85, 10.9, .46, 'Complete-JSON accuracy measures all fields together for the same sentence.', size=22, bold=True, align=PP_ALIGN.CENTER)
source_line(s, 'Recomputed from saved responses. Category: normalized text. Target: exact string. Label: strict integer match.')

# 19. Per-class failure pattern.
labels = ['Anti-stereotype', 'Stereotype', 'Neutral without target', 'Neutral with target', 'Bias']
ft_report = R['ft_labels']['classification_report']
json_report = R['ft_json']['classification_report']
modern_report = R['modernbert']['classification_report']
s = slide('Experimental Results', 'The main remaining errors separate stereotypes from anti-stereotypes', prototype=6,
          sources=[DATA['sources']['ft_labels'], DATA['sources']['ft_json'], DATA['sources']['modernbert']], speaker='The first two classes have lower F1 than the neutrality and bias classes in all three of these saved runs. In the fine-tuned label model, 25 anti-stereotype examples become stereotype predictions, and 48 stereotypes become anti-stereotype predictions: 73 of the 78 test mistakes are swaps between these classes. This is a count from the saved confusion matrix, not a linguistic causal explanation. The ModernBERT-large column uses the saved additional run, not the earlier review score.')
values = [[labels[i], str(DATA['class_support'][i]), f'{modern_report[str(i)]["f1-score"]*100:.2f}', f'{ft_report[labels[i] if i<2 or i==4 else ["Neutral without target term", "Neutral with target term"][i-2]]["f1-score"]*100:.2f}', f'{json_report[str(i)]["f1-score"]*100:.2f}'] for i in range(5)]
table(s, ['Gold class', 'Support', 'ModernBERT-large', 'Qwen label FT', 'Qwen JSON FT'], values,
      [3.45, 1.1, 2.4, 2.15, 2.15], y=2.12, row_height=.56, size=18)
box(s, 1.02, 5.78, 11.25, .56, PALE_AMBER)
text(s, 1.22, 5.88, 10.85, .35, '73 of 78 label-model errors are stereotype ↔ anti-stereotype swaps.', size=22, bold=True)
source_line(s, 'Table entries are per-class F1 (%). ModernBERT-large uses its saved run. FT = task fine-tuned.')

# 20. Demonstration: exact saved JSON outputs, benign example.
s = slide('Demonstration', 'Saved test example · Same sentence, different structured predictions', prototype=7,
          sources=[base_fields['source'], tuned_fields['source'], 'data/test.csv'], speaker='This demonstration replays saved responses for test index 0 and is not a fresh model inference. Dataset annotation: Category Neutral, Target Neutral, Label 2. The base model instead targets dog and predicts class 3. The task adapter recovers the annotated neutral target and class 2. This is a single illustration rather than an estimate of overall accuracy.')
box(s, 1.02, 2.13, 11.25, .7, PALE_BLUE)
text(s, 1.25, 2.3, 10.79, .38, 'Input sentence: “The dog begged for his treat.”', size=25, bold=True)
card(s, 1.02, 3.18, 5.46, 2.62, 'Base checkpoint', '{\n  "category": "neutral",\n  "target": "dog", "label": 3\n}', PALE_AMBER, 22)
card(s, 6.72, 3.18, 5.57, 2.62, 'Fine-tuned JSON adapter', '{\n  "category": "neutral",\n  "target": "Neutral", "label": 2\n}', PALE_GREEN, 22)
text(s, 1.2, 6.02, 10.88, .4, 'Gold annotation: category neutral · target Neutral · label 2', size=22, align=PP_ALIGN.CENTER)
source_line(s, 'Exact saved completions for test index 0. This example is a recorded demonstration.')

# 21. Demonstration: actual saved attributions, editable vector bars.
ex = DATA['explanation']
detail = ex['explanations']['1']
s = slide('Demonstration', 'ModernBERT-large · A saved LIME/SHAP explanation', prototype=7,
          sources=[DATA['explanation_source']], speaker='Exact saved diagnostic example: test index 1713, Producers are businesslike. Both gold label and prediction are class 1 (Stereotype). Bars use the saved LIME and SHAP scores for that predicted class. These are different explanation methods on the same word positions, not independent word probabilities. SHAP contribution for businesslike is 0.6305265 and LIME weight 0.6884328. Removing the top-ranked word drops the class score by 0.5773495. Raw softmax scores are uncalibrated; no correctness confidence is inferred. A hand-selected example is not a population benchmark.')
box(s, 1.02, 2.07, 11.25, .67, PALE_BLUE)
text(s, 1.22, 2.23, 10.85, .4, '“Producers are businesslike” · gold 1 · predicted 1 (Stereotype)', size=23, bold=True)
text(s, 3.1, 3.01, 3.0, .3, 'LIME', size=18, bold=True, color=BLUE)
text(s, 6.28, 3.01, 3.0, .3, 'SHAP', size=18, bold=True, color=GREEN)
for i, word in enumerate(ex['words']):
    y = 3.6+i*.66
    text(s, 1.08, y, 1.85, .4, word, size=22)
    for x, key, color in [(3.1, 'lime_values', '8FAAD8'), (6.28, 'shap_values', '81AC90')]:
        value = detail[key][i]
        width = value / .75 * 2.34
        shape = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y+.02), Inches(width), Inches(.32))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()
        text(s, x+width+.07, y, .7, .35, f'{value:+.3f}', size=16)
box(s, 9.65, 3.14, 2.61, 2.13, PALE_AMBER)
text(s, 9.82, 3.36, 2.27, 1.56, 'Masking the top word drops the class score by 57.73 percentage points.', size=21, bold=True)
text(s, 1.15, 5.77, 11.0, .48, 'Positive bars support class 1 in this example; they are not word-level probabilities.', size=20)
source_line(s, 'Saved diagnostic test index 1713. Bar values are preserved from the recorded explanation run.')

# 22. Conclusion.
s = slide('Conclusion', 'Task adaptation improves predictions; difficult class boundaries remain', prototype=8,
          sources=[DATA['sources']['base_labels'], DATA['sources']['ft_labels'], DATA['sources']['base_json'], DATA['sources']['ft_json'], DATA['explanation_source']], speaker='Summarize measured outcomes without implying an input-controlled ranking between label and JSON tasks. Qwen label accuracy is 95.51% with Category and Target supplied; JSON label accuracy is 93.61% from Sentence only. Complete-JSON accuracy is 92.87%. Fine-tuning substantially improves both tasks over the no-adapter checkpoint in these recorded runs. Future work should use matched engines, repeated seeds, calibration, error analysis of class 0/1, and external-domain evaluation. Word-level attribution is a diagnostic and does not certify label correctness.')
card(s, 1.02, 2.1, 5.46, 3.79, 'What the experiments show', '• Qwen label accuracy: 95.51%\n• JSON label accuracy: 93.61%\n• Complete-JSON accuracy: 92.87%\n\nFine-tuning improves both task setups over the base checkpoint.', PALE_GREEN, 23)
card(s, 6.72, 2.1, 5.57, 3.79, 'Next phase', '• Review class 0 / class 1 errors\n• Repeat with matched engines and seeds\n• Evaluate calibration and new domains\n• Extend explanation diagnostics', PALE_BLUE, 23)
source_line(s, 'All reported outcomes use recorded full-test evaluations; future work is proposed, not completed.')

# 23–24. References, numbered consistently throughout the deck.
for start, stop in ((0, 6), (6, len(REFS))):
    s = slide('References', f'IEEE-style bibliography · {1 if start == 0 else 2} / 2', prototype=9,
              sources=[link for _, link in REFS[start:stop]], speaker='Public model descriptions are separated from local experimental results. Local results are traceable to repository paths in the corresponding speaker notes. All public links were consulted on October 22, 2026.')
    y = 2.04
    for index in range(start, stop):
        entry, url = REFS[index]
        text(s, 1.02, y, .42, .32, f'[{index+1}]', size=17, color=BLUE, bold=True)
        shape = text(s, 1.58, y, 10.6, .68, entry, size=17)
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.hyperlink.address = url
        y += .72 if stop-start == 6 else .84
    source_line(s, 'Click a reference to open the primary paper or official model card.')

# Drop the untouched template slides only after all duplicates have been created.
for _ in PROTOTYPES:
    slide_id = PRS.slides._sldIdLst[0]
    PRS.part.drop_rel(slide_id.rId)
    PRS.slides._sldIdLst.remove(slide_id)

PRS.core_properties.title = 'StereoDetect — BTP Review'
PRS.core_properties.subject = 'Model comparison, task adaptation, structured prediction, and explanations'
PRS.core_properties.author = 'StereoDetect Project Team'
PRS.core_properties.keywords = 'StereoDetect, ALBERT, ModernBERT-large, Qwen, MiniLM, LIME, SHAP'
PRS.core_properties.comments = 'Built locally from the supplied institutional template and recorded experiment outputs.'
output = FOLDER / 'StereoDetect_BTP_Review.pptx'
PRS.save(output)

# Mechanical checks: reopen, retain institution, editable diagrams, and allowed scope.
check = Presentation(output)
assert len(check.slides) == 24
visible = '\n'.join(shape.text for s in check.slides for shape in s.shapes if shape.has_text_frame)
assert 'thinking' not in visible.casefold()
assert 'modernbert-base' not in visible.casefold()
assert '33.54' not in visible  # Values are native editable chart labels/data.
for s in list(check.slides)[1:]:
    assert any(shape.has_text_frame and 'Indian Institute of Information Technology Kottayam' in shape.text for shape in s.shapes)
for slide_info, s in zip(SLIDE_LOG, check.slides):
    for shape in s.shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= check.slide_width + Inches(.01), (slide_info, shape.name)
        assert shape.top + shape.height <= check.slide_height + Inches(.01), (slide_info, shape.name)
(HERE / 'slide_outline.json').write_text(json.dumps(SLIDE_LOG, indent=2)+'\n')
print(f'Saved {len(check.slides)} slides: {output}')
