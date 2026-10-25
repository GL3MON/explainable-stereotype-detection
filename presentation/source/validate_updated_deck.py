"""Render private QA contact sheets and check the exported presentation."""
from pathlib import Path
import json
import pymupdf as fitz
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1] / 'revised_20261022'
DATA = json.loads((Path(__file__).resolve().parent / 'deck_data_updated.json').read_text())
prs = Presentation(ROOT / 'StereoDetect_BTP_Review_Updated.pptx')

def check_report(table, report, labels):
    assert len(table.rows) == len(labels)+1
    for row, label in zip(list(table.rows)[1:], labels):
        value=report[str(label)]
        actual=[cell.text for cell in row.cells][1:]
        expected=[f'{value[k]*100:.2f}' for k in ('precision','recall','f1-score')]+[str(int(value['support']))]
        assert actual == expected, (label, actual, expected)

cases = iter([DATA['label_runs'][i] for i in (0,1,2,4,3,5,6,7,8)])
for number in range(29,34):
    for shape in prs.slides[number-1].shapes:
        if shape.has_table:
            check_report(shape.table,next(cases)['classification_report'],list(range(5)))
categories=['gender','race','profession','religion','neutral']
for number,field,indices,labels in [(34,'category',[0,2],categories),(35,'label',[0,2],list(range(5))),(36,'label',[1],list(range(5)))]:
    tables=[shape.table for shape in prs.slides[number-1].shapes if shape.has_table]
    for table,index in zip(tables,indices):
        check_report(table,DATA['json_runs'][index]['classification_reports'][field],labels)
for start,index in [(37,0),(40,2)]:
    covered=[]
    for number in range(start,start+3):
        for shape in prs.slides[number-1].shapes:
            if shape.has_table:
                labels=[row.cells[0].text for row in list(shape.table.rows)[1:]]
                check_report(shape.table,DATA['json_runs'][index]['classification_reports']['target'],labels)
                covered.extend(labels)
    assert covered==DATA['target_classes'] and len(covered)==87
for number,record in zip((23,24),DATA['explanation_examples']):
    all_text='\n'.join(s.text for s in prs.slides[number-1].shapes if s.has_text_frame)
    for key in ('lime_values','shap_values'):
        for value in record['explanations'][str(record['predicted_label'])][key]:
            assert f'{value:+.3f}' in all_text
QA = ROOT / 'qa'
QA.mkdir(exist_ok=True)
doc = fitz.open(ROOT / 'StereoDetect_BTP_Review_Updated.pdf')
assert len(doc) == 43, len(doc)
slides = []
for index, page in enumerate(doc, 1):
    pix = page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False)
    img = Image.frombytes('RGB', (pix.width, pix.height), pix.samples)
    img.save(QA / f'slide_{index:02d}.png')
    slides.append(img)
    assert page.get_text().strip(), index
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines', []):
            for span in line['spans']:
                x0, y0, x1, y1 = span['bbox']
                assert x0 >= -1 and y0 >= -1 and x1 <= page.rect.width + 1 and y1 <= page.rect.height + 1, (index, span['text'], span['bbox'])
for start in range(0, len(slides), 9):
    sheet = Image.new('RGB', (1440, 876), '#e4e8ee')
    draw = ImageDraw.Draw(sheet)
    for j, image in enumerate(slides[start:start+9]):
        image = image.copy()
        image.thumbnail((464, 261))
        x, y = (j % 3) * 480 + 8, (j // 3) * 292 + 22
        sheet.paste(image, (x, y))
        draw.text((x, y-17), f'Slide {start+j+1}', fill='black')
    sheet.save(QA / f'contact_sheet_{start//9+1}.png')
expected = {16:['23.19','93.84','70.66'], 20:['87'], 23:['businesslike','LIME','SHAP','0.70'], 24:['Gentlemen','car','LIME','SHAP','0.70']}
for number, snippets in expected.items():
    content = doc[number-1].get_text()
    for snippet in snippets:
        assert snippet in content, (number, snippet)
summary = {'slides':len(doc), 'pdf_text_bounds':'pass', 'key_text_checks':'pass', 'classification_table_values':'pass', 'target_coverage':'87/87 in both reports', 'signed_attribution_values':'pass', 'rendered_slides':len(slides)}
(QA/'validation.json').write_text(json.dumps(summary, indent=2)+'\n')
print(json.dumps(summary))
